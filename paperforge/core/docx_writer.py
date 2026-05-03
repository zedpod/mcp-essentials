"""DOCX rendering via python-docx. Writes the same markdown source to a .docx file.

Optional dependency. Returns False if `python-docx` is not installed; the caller
classifies that into a UNSUPPORTED Result error.
"""

import re
from pathlib import Path


def write_docx(markdown_source: str, *, output_path: Path, title: str) -> bool:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return False

    doc = Document()
    # Slightly better default body font.
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    lines = markdown_source.splitlines()
    # Strip YAML frontmatter — docx readers don't care.
    if lines and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                lines = lines[i + 1 :]
                break

    paragraph_buffer: list[str] = []

    def flush_paragraph():
        if not paragraph_buffer:
            return
        text = " ".join(paragraph_buffer).strip()
        if text:
            doc.add_paragraph(_strip_inline(text))
        paragraph_buffer.clear()

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            continue
        # Headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            flush_paragraph()
            level = len(m.group(1))
            doc.add_heading(_strip_inline(m.group(2)), level=min(level, 4))
            continue
        # Bullet
        if stripped.startswith(("- ", "* ")):
            flush_paragraph()
            doc.add_paragraph(_strip_inline(stripped[2:]), style="List Bullet")
            continue
        # Blockquote
        if stripped.startswith(">"):
            flush_paragraph()
            quote = doc.add_paragraph(_strip_inline(stripped[1:].strip()))
            quote.style = doc.styles["Intense Quote"] if "Intense Quote" in [s.name for s in doc.styles] else doc.styles["Normal"]
            continue
        # Horizontal rule
        if stripped == "---":
            flush_paragraph()
            doc.add_paragraph("─" * 40)
            continue
        paragraph_buffer.append(stripped)
    flush_paragraph()

    doc.save(str(output_path))
    return True


_INLINE_PATTERNS = [
    (re.compile(r"\*\*([^\*]+)\*\*"), r"\1"),
    (re.compile(r"(?<!\w)\*([^\*]+)\*(?!\w)"), r"\1"),
    (re.compile(r"`([^`]+)`"), r"\1"),
    (re.compile(r"\[([^\]]+)\]\([^)]+\)"), r"\1"),
]


def _strip_inline(text: str) -> str:
    """Remove markdown inline markers; docx body is plain runs.

    A richer renderer with bold/italic runs is left for v2 — most users edit
    the docx in Word and re-bold things to taste anyway.
    """
    out = text
    for pattern, repl in _INLINE_PATTERNS:
        out = pattern.sub(repl, out)
    return out
